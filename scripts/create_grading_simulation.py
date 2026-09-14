import sys, io, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

output_file = r'c:\Users\Newsk\Downloads\1ชทค\แบบจำลองการตัดเกรด_21909-1006_สัดส่วน_20-50-10-20.xlsx'

ROOM1_DATA = [
    (1, '69219090001', '001', 'นายกฤศมน บุญเปี่ยม', 66.0, 23, 24, 'เข้าสอบปกติ'),
    (2, '69219090002', '002', 'เด็กชายกษธณชล การรักเรียน', 58.0, 11, 18, 'เข้าสอบปกติ'),
    (3, '69219090003', '003', 'นางสาวกัญญ์วรา ทองดี', 87.0, 56, 44, 'เข้าสอบปกติ'),
    (4, '69219090004', '004', 'นายกันธรักษ์ พานทอง', 87.0, 20, 23, 'เข้าสอบปกติ'),
    (5, '69219090005', '005', 'เด็กชายคุณานนต์ พิมพ์จันทร์', 77.0, None, 32, 'ขาดสอบกลางภาค'),
    (6, '69219090006', '006', 'เด็กชายเฉลิมวงศ์ มะลิใย', 68.0, 23, 26, 'เข้าสอบปกติ'),
    (7, '69219090007', '007', 'นางสาวชลดา เรือนแก้ว', 35.0, 35, 23.5, 'เข้าสอบปกติ'),
    (8, '69219090008', '008', 'นายณัฐกรณ์ กองทอง', 43.0, None, None, 'ขาดสอบ'),
    (9, '69219090009', '009', 'นายณัฐพล ทวีทรัพย์', 76.0, 28, 21, 'เข้าสอบปกติ'),
    (10, '69219090010', '010', 'นายต้นตระการ เพชรภู่', 55.0, 13, 6, 'เข้าสอบปกติ'),
    (11, '69219090011', '011', 'นายตฤณ ตาดทอง', 88.0, 30, 30, 'เข้าสอบปกติ'),
    (12, '69219090012', '012', 'นายธนพัฒน์ บุญเรือง', 87.0, 50, 34, 'เข้าสอบปกติ'),
    (13, '69219090013', '013', 'นายธีรภัทร วาสิทธิ์เทพ', 78.0, 24, 22, 'เข้าสอบปกติ'),
    (14, '69219090014', '014', 'นายภูเบต ควันทอง', 9.0, None, None, 'ขาดสอบ'),
    (15, '69219090015', '015', 'นายภูวพิศิษฐ์ ใหลสกุล', 85.0, 71, 45, 'เข้าสอบปกติ'),
    (16, '69219090016', '016', 'นางสาวมินตา บินอามัด', 87.0, 34, 17, 'เข้าสอบปกติ'),
    (17, '69219090017', '017', 'นายอัรฮัม ยูนุ', 89.0, 35, 31, 'เข้าสอบปกติ'),
    (18, '69219090018', '018', 'นายอานนท์ณัฏฐ์ ไวพจน์', 84.0, 33, 26, 'เข้าสอบปกติ'),
]

ROOM2_DATA = [
    (1, '69219090020', '020', 'นายชุติวัต สุขมั่น', 78.0, 8, 16, 'เข้าสอบปกติ'),
    (2, '69219090021', '021', 'เด็กชายณัพร รัตนสุวรรณ', 38.0, 22, 15, 'เข้าสอบปกติ'),
    (3, '69219090022', '022', 'นายตุลยวัต ทาสีทอง', 53.0, 24, 35, 'เข้าสอบปกติ'),
    (4, '69219090023', '023', 'เด็กชายธนกฤต แก้วตา', 76.0, 31, 30, 'เข้าสอบปกติ'),
    (5, '69219090024', '024', 'นายธเนศพล กันเมือง', 60.0, 17, 23, 'เข้าสอบปกติ'),
    (6, '69219090025', '025', 'นายนรากร จันทร์กล่ำ', 29.0, None, None, 'ขาดสอบ'),
    (7, '69219090026', '026', 'นายนราวิชญ์ ทองระคนธ์', 56.0, 24, 25, 'เข้าสอบปกติ'),
    (8, '69219090027', '027', 'นายนันทพงศ์ ช่ออุบล', 15.0, 8, 8, 'เข้าสอบปกติ'),
    (9, '69219090028', '028', 'เด็กชายพงศ์พัทธ์ โสดา', 72.0, 24, 37, 'เข้าสอบปกติ'),
    (10, '69219090029', '029', 'นายภูธเรศ อินทโชติ', 87.0, 17, 31, 'เข้าสอบปกติ'),
    (11, '69219090030', '030', 'เด็กหญิงวราภรณ์ มั่นคง', 66.0, 9, 12, 'เข้าสอบปกติ'),
    (12, '69219090031', '031', 'นายวีรวัฒน์ วงษ์ทิพย์', 38.0, None, 13, 'ขาดสอบกลางภาค'),
    (13, '69219090032', '032', 'นายวุฒิกร พาณิชย์กุลเกียรติ', 82.0, 18, 14, 'เข้าสอบปกติ'),
    (14, '69219090033', '033', 'เด็กหญิงสโรชา มาเส็ง', 37.0, 14, 11, 'เข้าสอบปกติ'),
    (15, '69219090034', '034', 'นายอธิป ทองดอนเถื่อน', 47.0, 25, 29, 'เข้าสอบปกติ'),
    (16, '69219090035', '035', 'เด็กหญิงอารยา จันทน', 78.0, 24, 27, 'เข้าสอบปกติ'),
]

wb = openpyxl.Workbook()
wb.remove(wb.active)

font_main = 'Leelawadee UI'
title_font = Font(name=font_main, size=15, bold=True, color='0D47A1')
sub_font = Font(name=font_main, size=11, color='546E7A')
head_font = Font(name=font_main, size=10, bold=True, color='FFFFFF')
col_head_font = Font(name=font_main, size=10, bold=True, color='1A202C')
data_font = Font(name=font_main, size=10, color='2D3748')
bold_data_font = Font(name=font_main, size=10, bold=True, color='1A202C')
stat_label_font = Font(name=font_main, size=10, bold=True, color='1A365D')

fill_info = PatternFill("solid", fgColor="1A365D")
fill_aff = PatternFill("solid", fgColor="C2410C")
fill_assign = PatternFill("solid", fgColor="0F766E")
fill_mid = PatternFill("solid", fgColor="475569")
fill_fin = PatternFill("solid", fgColor="1D4ED8")
fill_grade = PatternFill("solid", fgColor="581C87")
fill_col_head = PatternFill("solid", fgColor="F1F5F9")
fill_zebra = PatternFill("solid", fgColor="F8FAFC")
fill_stat = PatternFill("solid", fgColor="E2E8F0")
fill_editable = PatternFill("solid", fgColor="FEF08A")
fill_grade_cell = PatternFill("solid", fgColor="F3E8FF")

thin_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)
double_bottom_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='double', color='1A365D')
)

def create_grading_sheet(sheet_title, room_label, data, default_mid_max):
    ws = wb.create_sheet(title=sheet_title)
    ws.views.sheetView[0].showGridLines = True
    
    ws.merge_cells('A1:O1')
    ws['A1'] = f"🎓 แบบจำลองการตัดเกรดผลการเรียนรู้ สัดส่วน 20 : 50 : 10 : 20 - {room_label}"
    ws['A1'].font = title_font
    ws['A1'].alignment = Alignment(horizontal='left', vertical='center')
    
    ws.merge_cells('A2:O2')
    ws['A2'] = f"รหัสวิชา 21909-1006 พื้นฐานการสร้างเว็บไซต์  |  สัดส่วนคะแนน: จิตพิสัย 20% + แบบฝึกหัด 50% + กลางภาค 10% + ปลายภาค 20% = 100%  |  {room_label}"
    ws['A2'].font = sub_font
    ws['A2'].alignment = Alignment(horizontal='left', vertical='center')
    
    super_groups = [
        ('A4:D4', 'ข้อมูลนักศึกษา', fill_info),
        ('E4:E4', 'จิตพิสัย (20%)', fill_aff),
        ('F4:G4', 'คะแนนเก็บ (50%)', fill_assign),
        ('H4:I4', 'กลางภาค (10%)', fill_mid),
        ('J4:K4', 'ปลายภาค (20%)', fill_fin),
        ('L4:O4', 'สรุปผลการเรียนและตัดเกรด (เต็ม 100)', fill_grade)
    ]
    for cell_rng, g_name, g_fill in super_groups:
        if ':' in cell_rng:
            ws.merge_cells(cell_rng)
            start_col, start_row = openpyxl.utils.coordinate_to_tuple(cell_rng.split(':')[0])
            end_col, end_row = openpyxl.utils.coordinate_to_tuple(cell_rng.split(':')[1])
            for r_idx in range(start_row, end_row + 1):
                for c_idx in range(start_col, end_col + 1):
                    c = ws.cell(row=r_idx, column=c_idx)
                    c.fill = g_fill
                    c.border = thin_border
            ws[cell_rng.split(':')[0]] = g_name
            ws[cell_rng.split(':')[0]].font = Font(name=font_main, size=11, bold=True, color='FFFFFF')
            ws[cell_rng.split(':')[0]].alignment = Alignment(horizontal='center', vertical='center')
        else:
            c = ws[cell_rng]
            c.value = g_name
            c.fill = g_fill
            c.font = Font(name=font_main, size=11, bold=True, color='FFFFFF')
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.border = thin_border
            
    sub_headers = [
        "เลขที่", "รหัสนักศึกษา", "รหัส\n(3 ตัว)", "ชื่อ - นามสกุล",
        "คะแนนจิตพิสัย\n(เต็ม 20)",
        "คะแนนดิบ\n(เต็ม 100)", "ทอนสัดส่วน\n(เต็ม 50)",
        "คะแนนดิบ\n(กลางภาค)", "ทอนสัดส่วน\n(เต็ม 10)",
        "คะแนนดิบ\n(เต็ม 50)", "ทอนสัดส่วน\n(เต็ม 20)",
        "คะแนนรวมสุทธิ\n(เต็ม 100)", "เกรด\n(0 - 4)", "ผลการเรียนรู้\n(การแปลผล)", "ลำดับที่\n(Rank)"
    ]
    ws.row_dimensions[4].height = 28
    ws.row_dimensions[5].height = 36
    ws.row_dimensions[6].height = 24
    
    for c_idx, h_text in enumerate(sub_headers, 1):
        c = ws.cell(row=5, column=c_idx, value=h_text)
        c.font = col_head_font
        c.fill = fill_col_head
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = thin_border
        
    max_vals = [
        None, None, None, "สัดส่วน / คะแนนเต็ม",
        20, 100, 50, default_mid_max, 10, 50, 20, 100, 4.0, "-", "-"
    ]
    for c_idx, val in enumerate(max_vals, 1):
        c = ws.cell(row=6, column=c_idx, value=val)
        c.font = bold_data_font
        c.fill = PatternFill("solid", fgColor="E0F2FE")
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = thin_border
        if c_idx == 4:
            c.alignment = Alignment(horizontal='right', vertical='center')
            
    start_r = 7
    last_r = start_r + len(data) - 1
    
    for idx, (no, full_id, code, name, as_raw, mid_raw, fin_raw, status) in enumerate(data):
        cur_r = start_r + idx
        ws.row_dimensions[cur_r].height = 22
        is_even = (idx % 2 == 1)
        row_fill = fill_zebra if is_even else PatternFill("solid", fgColor="FFFFFF")
        
        ws.cell(row=cur_r, column=1, value=no).alignment = Alignment(horizontal='center', vertical='center')
        ws.cell(row=cur_r, column=2, value=str(full_id)).alignment = Alignment(horizontal='center', vertical='center')
        ws.cell(row=cur_r, column=3, value=str(code)).alignment = Alignment(horizontal='center', vertical='center')
        ws.cell(row=cur_r, column=4, value=name).alignment = Alignment(horizontal='left', vertical='center')
        for c in range(1, 5):
            ws.cell(row=cur_r, column=c).font = data_font
            ws.cell(row=cur_r, column=c).fill = row_fill
            ws.cell(row=cur_r, column=c).border = thin_border
            
        default_aff = 20 if "ขาดสอบ" not in status else (10 if as_raw and as_raw > 0 else 0)
        c_aff = ws.cell(row=cur_r, column=5, value=default_aff)
        c_aff.font = Font(name=font_main, size=10, bold=True, color='854D0E')
        c_aff.fill = fill_editable
        c_aff.alignment = Alignment(horizontal='center', vertical='center')
        c_aff.border = thin_border
        c_aff.number_format = '0.0'
        
        c_as_raw = ws.cell(row=cur_r, column=6, value=as_raw)
        c_as_raw.font = data_font
        c_as_raw.fill = row_fill
        c_as_raw.alignment = Alignment(horizontal='center', vertical='center')
        c_as_raw.border = thin_border
        if as_raw is not None:
            c_as_raw.number_format = '0.0'
            
        c_as_w = ws.cell(row=cur_r, column=7)
        c_as_w.value = f'=IF(ISNUMBER(F{cur_r}), ROUND(F{cur_r}/$F$6*$G$6, 2), 0)'
        c_as_w.font = bold_data_font
        c_as_w.fill = PatternFill("solid", fgColor="CCFBF1")
        c_as_w.alignment = Alignment(horizontal='center', vertical='center')
        c_as_w.border = thin_border
        c_as_w.number_format = '0.00'
        
        c_mid_raw = ws.cell(row=cur_r, column=8, value=mid_raw if mid_raw is not None else "-")
        c_mid_raw.font = data_font
        c_mid_raw.fill = row_fill
        c_mid_raw.alignment = Alignment(horizontal='center', vertical='center')
        c_mid_raw.border = thin_border
        
        c_mid_w = ws.cell(row=cur_r, column=9)
        c_mid_w.value = f'=IF(ISNUMBER(H{cur_r}), ROUND(H{cur_r}/$H$6*$I$6, 2), 0)'
        c_mid_w.font = bold_data_font
        c_mid_w.fill = PatternFill("solid", fgColor="F1F5F9")
        c_mid_w.alignment = Alignment(horizontal='center', vertical='center')
        c_mid_w.border = thin_border
        c_mid_w.number_format = '0.00'
        
        c_fin_raw = ws.cell(row=cur_r, column=10, value=fin_raw if fin_raw is not None else "-")
        c_fin_raw.font = data_font
        c_fin_raw.fill = row_fill
        c_fin_raw.alignment = Alignment(horizontal='center', vertical='center')
        c_fin_raw.border = thin_border
        if isinstance(fin_raw, float):
            c_fin_raw.number_format = '0.0'
            
        c_fin_w = ws.cell(row=cur_r, column=11)
        c_fin_w.value = f'=IF(ISNUMBER(J{cur_r}), ROUND(J{cur_r}/$J$6*$K$6, 2), 0)'
        c_fin_w.font = bold_data_font
        c_fin_w.fill = PatternFill("solid", fgColor="DBEAFE")
        c_fin_w.alignment = Alignment(horizontal='center', vertical='center')
        c_fin_w.border = thin_border
        c_fin_w.number_format = '0.00'
        
        c_tot = ws.cell(row=cur_r, column=12)
        c_tot.value = f'=ROUND(E{cur_r}+G{cur_r}+I{cur_r}+K{cur_r}, 2)'
        c_tot.font = Font(name=font_main, size=11, bold=True, color='4A148C')
        c_tot.fill = fill_grade_cell
        c_tot.alignment = Alignment(horizontal='center', vertical='center')
        c_tot.border = thin_border
        c_tot.number_format = '0.00'
        
        c_grade = ws.cell(row=cur_r, column=13)
        c_grade.value = f'=IF(L{cur_r}>=80, 4, IF(L{cur_r}>=75, 3.5, IF(L{cur_r}>=70, 3, IF(L{cur_r}>=65, 2.5, IF(L{cur_r}>=60, 2, IF(L{cur_r}>=55, 1.5, IF(L{cur_r}>=50, 1, 0)))))))'
        c_grade.font = Font(name=font_main, size=11, bold=True, color='1A202C')
        c_grade.fill = fill_grade_cell
        c_grade.alignment = Alignment(horizontal='center', vertical='center')
        c_grade.border = thin_border
        c_grade.number_format = '0.0'
        
        c_eval = ws.cell(row=cur_r, column=14)
        c_eval.value = f'=IF(M{cur_r}>=4, "ดีเยี่ยม", IF(M{cur_r}>=3.5, "ดีมาก", IF(M{cur_r}>=3, "ดี", IF(M{cur_r}>=2.5, "ค่อนข้างดี", IF(M{cur_r}>=2, "พอใช้", IF(M{cur_r}>=1.5, "อ่อน", IF(M{cur_r}>=1, "อ่อนมาก", "ไม่ผ่าน")))))))'
        c_eval.font = data_font
        c_eval.fill = row_fill
        c_eval.alignment = Alignment(horizontal='center', vertical='center')
        c_eval.border = thin_border
        
        c_rk = ws.cell(row=cur_r, column=15)
        c_rk.value = f'=RANK(L{cur_r}, $L$7:$L${last_r})'
        c_rk.font = bold_data_font
        c_rk.fill = row_fill
        c_rk.alignment = Alignment(horizontal='center', vertical='center')
        c_rk.border = thin_border

    stat_rows_def = [
        ("คะแนนเฉลี่ย (Average)", "AVERAGE", True),
        ("คะแนนสูงสุด (Max)", "MAX", False),
        ("คะแนนต่ำสุด (Min)", "MIN", False),
        ("ส่วนเบี่ยงเบนมาตรฐาน (S.D.)", "STDEV.P", True),
    ]
    
    for s_idx, (s_label, s_func, is_round) in enumerate(stat_rows_def):
        r_num = last_r + 1 + s_idx
        ws.row_dimensions[r_num].height = 22
        ws.merge_cells(start_row=r_num, start_column=1, end_row=r_num, end_column=4)
        lbl_cell = ws.cell(row=r_num, column=1, value=s_label)
        lbl_cell.font = stat_label_font
        lbl_cell.alignment = Alignment(horizontal='right', vertical='center')
        
        for c in range(1, 5):
            ws.cell(row=r_num, column=c).fill = fill_stat
            ws.cell(row=r_num, column=c).border = thin_border if s_idx < 3 else double_bottom_border
            
        for col_idx in [5, 6, 7, 8, 9, 10, 11, 12, 13]:
            c_letter = get_column_letter(col_idx)
            c = ws.cell(row=r_num, column=col_idx)
            if is_round:
                c.value = f'=IFERROR(ROUND({s_func}({c_letter}7:{c_letter}{last_r}), 2), "-")'
            else:
                c.value = f'=IFERROR({s_func}({c_letter}7:{c_letter}{last_r}), "-")'
            c.font = bold_data_font
            c.fill = fill_stat
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.border = thin_border if s_idx < 3 else double_bottom_border
            c.number_format = '0.00'
            
        for c_idx in [14, 15]:
            c = ws.cell(row=r_num, column=c_idx)
            c.fill = fill_stat
            c.border = thin_border if s_idx < 3 else double_bottom_border

    widths = {1: 7, 2: 15, 3: 8, 4: 26, 5: 14, 6: 12, 7: 13, 8: 13, 9: 13, 10: 12, 11: 13, 12: 15, 13: 10, 14: 14, 15: 11}
    for c_idx, w in widths.items():
        ws.column_dimensions[get_column_letter(c_idx)].width = w

create_grading_sheet('ตัดเกรด ห้อง 1', 'ห้อง 1 (18 คน)', ROOM1_DATA, 100)
create_grading_sheet('ตัดเกรด ห้อง 2', 'ห้อง 2 (16 คน)', ROOM2_DATA, 50)

# Create 34-student master sheet
def create_all_students_sheet():
    ws = wb.create_sheet(title='ตัดเกรด รวม 2 ห้อง (34 คน)')
    ws.views.sheetView[0].showGridLines = True
    
    ws.merge_cells('A1:P1')
    ws['A1'] = '🎓 แบบจำลองการตัดเกรดผลการเรียนรู้ สัดส่วน 20 : 50 : 10 : 20 - รวมทั้ง 2 ห้อง (34 คน)'
    ws['A1'].font = title_font
    ws['A1'].alignment = Alignment(horizontal='left', vertical='center')
    
    ws.merge_cells('A2:P2')
    ws['A2'] = 'รหัสวิชา 21909-1006 พื้นฐานการสร้างเว็บไซต์  |  ระดับชั้น ปวช. 1 ชทค. (ห้อง 1 และ 2 รวม 34 คน)  |  ปีการศึกษา 2569'
    ws['A2'].font = sub_font
    ws['A2'].alignment = Alignment(horizontal='left', vertical='center')
    
    super_groups = [
        ('A4:E4', 'ข้อมูลนักศึกษา', fill_info),
        ('F4:F4', 'จิตพิสัย (20%)', fill_aff),
        ('G4:H4', 'คะแนนเก็บ (50%)', fill_assign),
        ('I4:J4', 'กลางภาค (10%)', fill_mid),
        ('K4:L4', 'ปลายภาค (20%)', fill_fin),
        ('M4:P4', 'สรุปผลการเรียนและตัดเกรด (เต็ม 100)', fill_grade)
    ]
    for cell_rng, g_name, g_fill in super_groups:
        if ':' in cell_rng:
            ws.merge_cells(cell_rng)
            start_col, start_row = openpyxl.utils.coordinate_to_tuple(cell_rng.split(':')[0])
            end_col, end_row = openpyxl.utils.coordinate_to_tuple(cell_rng.split(':')[1])
            for r_idx in range(start_row, end_row + 1):
                for c_idx in range(start_col, end_col + 1):
                    c = ws.cell(row=r_idx, column=c_idx)
                    c.fill = g_fill
                    c.border = thin_border
            ws[cell_rng.split(':')[0]] = g_name
            ws[cell_rng.split(':')[0]].font = Font(name=font_main, size=11, bold=True, color='FFFFFF')
            ws[cell_rng.split(':')[0]].alignment = Alignment(horizontal='center', vertical='center')
        else:
            c = ws[cell_rng]
            c.value = g_name
            c.fill = g_fill
            c.font = Font(name=font_main, size=11, bold=True, color='FFFFFF')
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.border = thin_border

    sub_headers = [
        'ลำดับรวม', 'ห้อง', 'เลขที่', 'รหัส', 'ชื่อ - นามสกุล',
        'จิตพิสัย\n(20)', 'ดิบเก็บ\n(100)', 'ทอนเก็บ\n(50)',
        'ดิบกลาง', 'ทอนกลาง\n(10)', 'ดิบปลาย\n(50)', 'ทอนปลาย\n(20)',
        'คะแนนสุทธิ\n(100)', 'เกรด\n(0-4)', 'ผลการเรียนรู้', 'ลำดับรวม\n(Rank)'
    ]
    ws.row_dimensions[4].height = 28
    ws.row_dimensions[5].height = 36
    ws.row_dimensions[6].height = 24

    for c_idx, h_text in enumerate(sub_headers, 1):
        c = ws.cell(row=5, column=c_idx, value=h_text)
        c.font = col_head_font
        c.fill = fill_col_head
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = thin_border

    max_vals = [None, None, None, None, 'สัดส่วน / คะแนนเต็ม', 20, 100, 50, 100, 10, 50, 20, 100, 4.0, '-', '-']
    for c_idx, val in enumerate(max_vals, 1):
        c = ws.cell(row=6, column=c_idx, value=val)
        c.font = bold_data_font
        c.fill = PatternFill('solid', fgColor='E0F2FE')
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = thin_border
        if c_idx == 5:
            c.alignment = Alignment(horizontal='right', vertical='center')

    cur_r = 7
    # Room 1 rows
    for r in range(7, 25):
        ws.row_dimensions[cur_r].height = 22
        row_fill = fill_zebra if cur_r % 2 == 1 else PatternFill('solid', fgColor='FFFFFF')
        
        ws.cell(row=cur_r, column=1, value=cur_r-6).alignment = Alignment(horizontal='center', vertical='center')
        ws.cell(row=cur_r, column=2, value='ห้อง 1').alignment = Alignment(horizontal='center', vertical='center')
        ws.cell(row=cur_r, column=3, value=f"='ตัดเกรด ห้อง 1'!A{r}").alignment = Alignment(horizontal='center', vertical='center')
        ws.cell(row=cur_r, column=4, value=f"='ตัดเกรด ห้อง 1'!C{r}").alignment = Alignment(horizontal='center', vertical='center')
        ws.cell(row=cur_r, column=5, value=f"='ตัดเกรด ห้อง 1'!D{r}").alignment = Alignment(horizontal='left', vertical='center')
        
        for c in range(1, 6):
            ws.cell(row=cur_r, column=c).font = data_font
            ws.cell(row=cur_r, column=c).fill = row_fill
            ws.cell(row=cur_r, column=c).border = thin_border
            
        ws.cell(row=cur_r, column=6, value=f"='ตัดเกรด ห้อง 1'!E{r}").number_format = '0.0'
        ws.cell(row=cur_r, column=7, value=f"='ตัดเกรด ห้อง 1'!F{r}").number_format = '0.0'
        ws.cell(row=cur_r, column=8, value=f"='ตัดเกรด ห้อง 1'!G{r}").number_format = '0.00'
        ws.cell(row=cur_r, column=9, value=f"='ตัดเกรด ห้อง 1'!H{r}")
        ws.cell(row=cur_r, column=10, value=f"='ตัดเกรด ห้อง 1'!I{r}").number_format = '0.00'
        ws.cell(row=cur_r, column=11, value=f"='ตัดเกรด ห้อง 1'!J{r}")
        ws.cell(row=cur_r, column=12, value=f"='ตัดเกรด ห้อง 1'!K{r}").number_format = '0.00'
        
        for c in range(6, 13):
            ws.cell(row=cur_r, column=c).font = data_font
            ws.cell(row=cur_r, column=c).fill = row_fill
            ws.cell(row=cur_r, column=c).alignment = Alignment(horizontal='center', vertical='center')
            ws.cell(row=cur_r, column=c).border = thin_border
            
        c_tot = ws.cell(row=cur_r, column=13, value=f"='ตัดเกรด ห้อง 1'!L{r}")
        c_tot.font = Font(name=font_main, size=11, bold=True, color='4A148C')
        c_tot.fill = fill_grade_cell
        c_tot.alignment = Alignment(horizontal='center', vertical='center')
        c_tot.border = thin_border
        c_tot.number_format = '0.00'
        
        c_gr = ws.cell(row=cur_r, column=14, value=f"='ตัดเกรด ห้อง 1'!M{r}")
        c_gr.font = Font(name=font_main, size=11, bold=True, color='1A202C')
        c_gr.fill = fill_grade_cell
        c_gr.alignment = Alignment(horizontal='center', vertical='center')
        c_gr.border = thin_border
        c_gr.number_format = '0.0'
        
        c_ev = ws.cell(row=cur_r, column=15, value=f"='ตัดเกรด ห้อง 1'!N{r}")
        c_ev.font = data_font
        c_ev.fill = row_fill
        c_ev.alignment = Alignment(horizontal='center', vertical='center')
        c_ev.border = thin_border
        
        c_rk = ws.cell(row=cur_r, column=16, value=f'=RANK(M{cur_r}, $M$7:$M$40)')
        c_rk.font = bold_data_font
        c_rk.fill = row_fill
        c_rk.alignment = Alignment(horizontal='center', vertical='center')
        c_rk.border = thin_border
        
        cur_r += 1

    # Room 2 rows
    for r in range(7, 23):
        ws.row_dimensions[cur_r].height = 22
        row_fill = fill_zebra if cur_r % 2 == 1 else PatternFill('solid', fgColor='FFFFFF')
        
        ws.cell(row=cur_r, column=1, value=cur_r-6).alignment = Alignment(horizontal='center', vertical='center')
        ws.cell(row=cur_r, column=2, value='ห้อง 2').alignment = Alignment(horizontal='center', vertical='center')
        ws.cell(row=cur_r, column=3, value=f"='ตัดเกรด ห้อง 2'!A{r}").alignment = Alignment(horizontal='center', vertical='center')
        ws.cell(row=cur_r, column=4, value=f"='ตัดเกรด ห้อง 2'!C{r}").alignment = Alignment(horizontal='center', vertical='center')
        ws.cell(row=cur_r, column=5, value=f"='ตัดเกรด ห้อง 2'!D{r}").alignment = Alignment(horizontal='left', vertical='center')
        
        for c in range(1, 6):
            ws.cell(row=cur_r, column=c).font = data_font
            ws.cell(row=cur_r, column=c).fill = row_fill
            ws.cell(row=cur_r, column=c).border = thin_border
            
        ws.cell(row=cur_r, column=6, value=f"='ตัดเกรด ห้อง 2'!E{r}").number_format = '0.0'
        ws.cell(row=cur_r, column=7, value=f"='ตัดเกรด ห้อง 2'!F{r}").number_format = '0.0'
        ws.cell(row=cur_r, column=8, value=f"='ตัดเกรด ห้อง 2'!G{r}").number_format = '0.00'
        ws.cell(row=cur_r, column=9, value=f"='ตัดเกรด ห้อง 2'!H{r}")
        ws.cell(row=cur_r, column=10, value=f"='ตัดเกรด ห้อง 2'!I{r}").number_format = '0.00'
        ws.cell(row=cur_r, column=11, value=f"='ตัดเกรด ห้อง 2'!J{r}")
        ws.cell(row=cur_r, column=12, value=f"='ตัดเกรด ห้อง 2'!K{r}").number_format = '0.00'
        
        for c in range(6, 13):
            ws.cell(row=cur_r, column=c).font = data_font
            ws.cell(row=cur_r, column=c).fill = row_fill
            ws.cell(row=cur_r, column=c).alignment = Alignment(horizontal='center', vertical='center')
            ws.cell(row=cur_r, column=c).border = thin_border
            
        c_tot = ws.cell(row=cur_r, column=13, value=f"='ตัดเกรด ห้อง 2'!L{r}")
        c_tot.font = Font(name=font_main, size=11, bold=True, color='4A148C')
        c_tot.fill = fill_grade_cell
        c_tot.alignment = Alignment(horizontal='center', vertical='center')
        c_tot.border = thin_border
        c_tot.number_format = '0.00'
        
        c_gr = ws.cell(row=cur_r, column=14, value=f"='ตัดเกรด ห้อง 2'!M{r}")
        c_gr.font = Font(name=font_main, size=11, bold=True, color='1A202C')
        c_gr.fill = fill_grade_cell
        c_gr.alignment = Alignment(horizontal='center', vertical='center')
        c_gr.border = thin_border
        c_gr.number_format = '0.0'
        
        c_ev = ws.cell(row=cur_r, column=15, value=f"='ตัดเกรด ห้อง 2'!N{r}")
        c_ev.font = data_font
        c_ev.fill = row_fill
        c_ev.alignment = Alignment(horizontal='center', vertical='center')
        c_ev.border = thin_border
        
        c_rk = ws.cell(row=cur_r, column=16, value=f'=RANK(M{cur_r}, $M$7:$M$40)')
        c_rk.font = bold_data_font
        c_rk.fill = row_fill
        c_rk.alignment = Alignment(horizontal='center', vertical='center')
        c_rk.border = thin_border
        
        cur_r += 1

    last_r = 40
    stat_rows_def = [
        ("คะแนนเฉลี่ย (Average)", "AVERAGE", True),
        ("คะแนนสูงสุด (Max)", "MAX", False),
        ("คะแนนต่ำสุด (Min)", "MIN", False),
        ("ส่วนเบี่ยงเบนมาตรฐาน (S.D.)", "STDEV.P", True),
    ]
    for s_idx, (s_label, s_func, is_round) in enumerate(stat_rows_def):
        r_num = last_r + 1 + s_idx
        ws.row_dimensions[r_num].height = 22
        ws.merge_cells(start_row=r_num, start_column=1, end_row=r_num, end_column=5)
        lbl_cell = ws.cell(row=r_num, column=1, value=s_label)
        lbl_cell.font = stat_label_font
        lbl_cell.alignment = Alignment(horizontal='right', vertical='center')
        
        for c in range(1, 6):
            ws.cell(row=r_num, column=c).fill = fill_stat
            ws.cell(row=r_num, column=c).border = thin_border if s_idx < 3 else double_bottom_border
            
        for col_idx in [6, 7, 8, 9, 10, 11, 12, 13, 14]:
            c_letter = get_column_letter(col_idx)
            c = ws.cell(row=r_num, column=col_idx)
            if is_round:
                c.value = f'=IFERROR(ROUND({s_func}({c_letter}7:{c_letter}{last_r}), 2), "-")'
            else:
                c.value = f'=IFERROR({s_func}({c_letter}7:{c_letter}{last_r}), "-")'
            c.font = bold_data_font
            c.fill = fill_stat
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.border = thin_border if s_idx < 3 else double_bottom_border
            c.number_format = '0.00'
            
        for c_idx in [15, 16]:
            c = ws.cell(row=r_num, column=c_idx)
            c.fill = fill_stat
            c.border = thin_border if s_idx < 3 else double_bottom_border

    col_w = {1: 8, 2: 9, 3: 7, 4: 8, 5: 26, 6: 12, 7: 12, 8: 12, 9: 12, 10: 12, 11: 12, 12: 12, 13: 14, 14: 10, 15: 14, 16: 11}
    for c_idx, w in col_w.items():
        ws.column_dimensions[get_column_letter(c_idx)].width = w

create_all_students_sheet()

# Create Overview Sheet
def create_overview_sheet():
    ws = wb.create_sheet(title='สรุปภาพรวมและสถิติเกรด', index=0)
    ws.views.sheetView[0].showGridLines = True
    
    ws.merge_cells('A1:F1')
    ws['A1'] = "📊 รายงานสรุปผลการประเมินและแจกแจงเกรด สัดส่วน 20 : 50 : 10 : 20"
    ws['A1'].font = title_font
    ws['A1'].alignment = Alignment(horizontal='left', vertical='center')
    
    ws.merge_cells('A2:F2')
    ws['A2'] = "รหัสวิชา 21909-1006 พื้นฐานการสร้างเว็บไซต์  |  ระดับชั้น ปวช. 1 ชทค. (ห้อง 1 และ 2 รวม 34 คน)  |  ปีการศึกษา 2569"
    ws['A2'].font = sub_font
    ws['A2'].alignment = Alignment(horizontal='left', vertical='center')
    
    ws.cell(row=4, column=1, value="📌 โครงสร้างสัดส่วนการตัดเกรด 100 คะแนน").font = Font(name=font_main, size=12, bold=True, color="1A365D")
    
    struct_headers = ["หมวดการประเมิน", "สัดส่วน", "คะแนนเต็ม", "วิธีการแปลงคะแนนจากข้อมูลจริง", "เกณฑ์อ้างอิง"]
    for c_idx, h in enumerate(struct_headers, 1):
        c = ws.cell(row=5, column=c_idx, value=h)
        c.font = head_font
        c.fill = fill_info
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = thin_border
        
    struct_data = [
        ("1. จิตพิสัย (คุณธรรม จริยธรรม เจตคติ)", "20%", "20 คะแนน", "ตั้งค่าเริ่มต้น 20 คะแนน (แก้ไขปรับลดตาม ขาด/ลา/สาย ได้)", "ระเบียบ สอศ. ข้อ 4"),
        ("2. คะแนนเก็บระหว่างภาค (แบบฝึกหัด 1-8 + โปรเจค)", "50%", "50 คะแนน", "ทอนจากคะแนนเก็บเต็ม 100 คะแนน = (คะแนนดิบ ÷ 100) × 50", "งานใน Google Classroom"),
        ("3. สอบกลางภาค (Midterm Examination)", "10%", "10 คะแนน", "ทอนจากคะแนนสอบกลางภาค = (คะแนนดิบ ÷ คะแนนเต็มดิบ) × 10", "ข้อสอบกลางภาค"),
        ("4. สอบปลายภาค (Final Examination)", "20%", "20 คะแนน", "ทอนจากคะแนนสอบปลายภาคเต็ม 50 = (คะแนนดิบ ÷ 50) × 20", "ปรนัย 30 + อัตนัย 20"),
        ("รวมทั้งสิ้น (Total Net Score)", "100%", "100 คะแนน", "รวมทุกหมวดเข้าด้วยกัน ตัดเกรด 8 ระดับ (0 - 4.0)", "เกณฑ์มาตรฐาน ปวช.")
    ]
    for r_idx, row_vals in enumerate(struct_data, 6):
        for c_idx, val in enumerate(row_vals, 1):
            c = ws.cell(row=r_idx, column=c_idx, value=val)
            c.font = bold_data_font if r_idx == 10 else data_font
            c.border = thin_border
            if r_idx == 10:
                c.fill = PatternFill("solid", fgColor="E0F2FE")
            if c_idx in [2, 3]:
                c.alignment = Alignment(horizontal='center', vertical='center')
            else:
                c.alignment = Alignment(horizontal='left', vertical='center')
                
    start_r = 13
    ws.cell(row=start_r, column=1, value="📈 การแจกแจงระดับผลการเรียน (Grade Distribution 8 ระดับ)").font = Font(name=font_main, size=12, bold=True, color="1A365D")
    
    dist_headers = ["ช่วงคะแนน (100)", "เกรด", "ความหมาย", "ห้อง 1 (18 คน)", "ห้อง 2 (16 คน)", "รวม (34 คน)", "ร้อยละ"]
    for c_idx, h in enumerate(dist_headers, 1):
        c = ws.cell(row=start_r+1, column=c_idx, value=h)
        c.font = head_font
        c.fill = fill_grade
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = thin_border
        
    grade_rows = [
        ("80 - 100 คะแนน", "4.0", "ดีเยี่ยม (Excellent)",
         "=COUNTIFS('ตัดเกรด ห้อง 1'!L7:L24, \">=80\", 'ตัดเกรด ห้อง 1'!L7:L24, \"<=100\")",
         "=COUNTIFS('ตัดเกรด ห้อง 2'!L7:L22, \">=80\", 'ตัดเกรด ห้อง 2'!L7:L22, \"<=100\")"),
        ("75 - 79 คะแนน", "3.5", "ดีมาก (Very Good)",
         "=COUNTIFS('ตัดเกรด ห้อง 1'!L7:L24, \">=75\", 'ตัดเกรด ห้อง 1'!L7:L24, \"<80\")",
         "=COUNTIFS('ตัดเกรด ห้อง 2'!L7:L22, \">=75\", 'ตัดเกรด ห้อง 2'!L7:L22, \"<80\")"),
        ("70 - 74 คะแนน", "3.0", "ดี (Good)",
         "=COUNTIFS('ตัดเกรด ห้อง 1'!L7:L24, \">=70\", 'ตัดเกรด ห้อง 1'!L7:L24, \"<75\")",
         "=COUNTIFS('ตัดเกรด ห้อง 2'!L7:L22, \">=70\", 'ตัดเกรด ห้อง 2'!L7:L22, \"<75\")"),
        ("65 - 69 คะแนน", "2.5", "ค่อนข้างดี (Fairly Good)",
         "=COUNTIFS('ตัดเกรด ห้อง 1'!L7:L24, \">=65\", 'ตัดเกรด ห้อง 1'!L7:L24, \"<70\")",
         "=COUNTIFS('ตัดเกรด ห้อง 2'!L7:L22, \">=65\", 'ตัดเกรด ห้อง 2'!L7:L22, \"<70\")"),
        ("60 - 64 คะแนน", "2.0", "พอใช้ (Fair)",
         "=COUNTIFS('ตัดเกรด ห้อง 1'!L7:L24, \">=60\", 'ตัดเกรด ห้อง 1'!L7:L24, \"<65\")",
         "=COUNTIFS('ตัดเกรด ห้อง 2'!L7:L22, \">=60\", 'ตัดเกรด ห้อง 2'!L7:L22, \"<65\")"),
        ("55 - 59 คะแนน", "1.5", "อ่อน (Poor)",
         "=COUNTIFS('ตัดเกรด ห้อง 1'!L7:L24, \">=55\", 'ตัดเกรด ห้อง 1'!L7:L24, \"<60\")",
         "=COUNTIFS('ตัดเกรด ห้อง 2'!L7:L22, \">=55\", 'ตัดเกรด ห้อง 2'!L7:L22, \"<60\")"),
        ("50 - 54 คะแนน", "1.0", "อ่อนมาก (Very Poor)",
         "=COUNTIFS('ตัดเกรด ห้อง 1'!L7:L24, \">=50\", 'ตัดเกรด ห้อง 1'!L7:L24, \"<55\")",
         "=COUNTIFS('ตัดเกรด ห้อง 2'!L7:L22, \">=50\", 'ตัดเกรด ห้อง 2'!L7:L22, \"<55\")"),
        ("0 - 49 คะแนน", "0.0", "ไม่ผ่านเกณฑ์ (Fail)",
         "=COUNTIFS('ตัดเกรด ห้อง 1'!L7:L24, \">=0\", 'ตัดเกรด ห้อง 1'!L7:L24, \"<50\")",
         "=COUNTIFS('ตัดเกรด ห้อง 2'!L7:L22, \">=0\", 'ตัดเกรด ห้อง 2'!L7:L22, \"<50\")"),
    ]
    
    for r_offset, (rng, gr, desc, f1, f2) in enumerate(grade_rows, 2):
        cur_r = start_r + r_offset
        ws.cell(row=cur_r, column=1, value=rng).alignment = Alignment(horizontal='center', vertical='center')
        ws.cell(row=cur_r, column=2, value=gr).alignment = Alignment(horizontal='center', vertical='center')
        ws.cell(row=cur_r, column=3, value=desc).alignment = Alignment(horizontal='left', vertical='center')
        ws.cell(row=cur_r, column=4, value=f1).alignment = Alignment(horizontal='center', vertical='center')
        ws.cell(row=cur_r, column=5, value=f2).alignment = Alignment(horizontal='center', vertical='center')
        ws.cell(row=cur_r, column=6, value=f"=SUM(D{cur_r}:E{cur_r})").alignment = Alignment(horizontal='center', vertical='center')
        ws.cell(row=cur_r, column=7, value=f"=ROUND(F{cur_r}/34*100, 1) & \"%\"").alignment = Alignment(horizontal='center', vertical='center')
        for c in range(1, 8):
            cell = ws.cell(row=cur_r, column=c)
            cell.font = data_font
            cell.border = thin_border
            if gr in ["4.0", "3.5"]:
                cell.fill = PatternFill("solid", fgColor="F0FDF4")
            elif gr in ["0.0"]:
                cell.fill = PatternFill("solid", fgColor="FEF2F2")
                
    tot_r = start_r + len(grade_rows) + 2
    ws.cell(row=tot_r, column=1, value="รวมนักเรียนทั้งหมด").alignment = Alignment(horizontal='center', vertical='center')
    ws.cell(row=tot_r, column=2, value="-").alignment = Alignment(horizontal='center', vertical='center')
    ws.cell(row=tot_r, column=3, value="34 คน").alignment = Alignment(horizontal='left', vertical='center')
    ws.cell(row=tot_r, column=4, value=f"=SUM(D{start_r+2}:D{tot_r-1})").alignment = Alignment(horizontal='center', vertical='center')
    ws.cell(row=tot_r, column=5, value=f"=SUM(E{start_r+2}:E{tot_r-1})").alignment = Alignment(horizontal='center', vertical='center')
    ws.cell(row=tot_r, column=6, value=f"=SUM(F{start_r+2}:F{tot_r-1})").alignment = Alignment(horizontal='center', vertical='center')
    ws.cell(row=tot_r, column=7, value="100.0%").alignment = Alignment(horizontal='center', vertical='center')
    for c in range(1, 8):
        cell = ws.cell(row=tot_r, column=c)
        cell.font = bold_data_font
        cell.fill = fill_stat
        cell.border = double_bottom_border

    m_start = tot_r + 3
    ws.cell(row=m_start, column=1, value="🏆 สรุปตัวชี้วัดผลสัมฤทธิ์ทางการเรียน (Key Academic Metrics)").font = Font(name=font_main, size=12, bold=True, color="1A365D")
    
    m_headers = ["ตัวชี้วัด", "ห้อง 1", "ห้อง 2", "รวมทั้งสิ้น / ค่าเฉลี่ยรวม", "หมายเหตุ"]
    for c_idx, h in enumerate(m_headers, 1):
        c = ws.cell(row=m_start+1, column=c_idx, value=h)
        c.font = head_font
        c.fill = fill_info
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = thin_border
        
    metrics = [
        ("เกรดเฉลี่ยรายวิชา (GPA)", 
         "=ROUND(AVERAGE('ตัดเกรด ห้อง 1'!M7:M24), 2)", 
         "=ROUND(AVERAGE('ตัดเกรด ห้อง 2'!M7:M22), 2)", 
         "=ROUND((AVERAGE('ตัดเกรด ห้อง 1'!M7:M24)*18 + AVERAGE('ตัดเกรด ห้อง 2'!M7:M22)*16)/34, 2)", 
         "เกรดเฉลี่ยประชากรนักเรียน"),
        ("คะแนนเฉลี่ยรวม (เต็ม 100)", 
         "=ROUND(AVERAGE('ตัดเกรด ห้อง 1'!L7:L24), 2)", 
         "=ROUND(AVERAGE('ตัดเกรด ห้อง 2'!L7:L22), 2)", 
         "=ROUND((AVERAGE('ตัดเกรด ห้อง 1'!L7:L24)*18 + AVERAGE('ตัดเกรด ห้อง 2'!L7:L22)*16)/34, 2)", 
         "รวมทุกสัดส่วน 100 คะแนน"),
        ("คะแนนสูงสุด (Max)", 
         "=MAX('ตัดเกรด ห้อง 1'!L7:L24)", 
         "=MAX('ตัดเกรด ห้อง 2'!L7:L22)", 
         "=MAX('ตัดเกรด ห้อง 1'!L7:L24, 'ตัดเกรด ห้อง 2'!L7:L22)", 
         "คะแนนรวมสูงสุดในแต่ละห้อง"),
        ("คะแนนต่ำสุด (Min)", 
         "=MIN('ตัดเกรด ห้อง 1'!L7:L24)", 
         "=MIN('ตัดเกรด ห้อง 2'!L7:L22)", 
         "=MIN('ตัดเกรด ห้อง 1'!L7:L24, 'ตัดเกรด ห้อง 2'!L7:L22)", 
         "คะแนนรวมต่ำสุดในแต่ละห้อง"),
        ("จำนวนนักเรียนที่ผ่านเกณฑ์ (เกรด ≥ 1.0)", 
         "=COUNTIF('ตัดเกรด ห้อง 1'!M7:M24, \">=1.0\")", 
         "=COUNTIF('ตัดเกรด ห้อง 2'!M7:M22, \">=1.0\")", 
         "=COUNTIF('ตัดเกรด ห้อง 1'!M7:M24, \">=1.0\") + COUNTIF('ตัดเกรด ห้อง 2'!M7:M22, \">=1.0\")", 
         "เกณฑ์ผ่านร้อยละ 50 ของคะแนนเต็ม"),
        ("ร้อยละการผ่านเกณฑ์", 
         f"=ROUND(B{m_start+6}/18*100, 1) & \"%\"", 
         f"=ROUND(C{m_start+6}/16*100, 1) & \"%\"", 
         f"=ROUND(D{m_start+6}/34*100, 1) & \"%\"", 
         "เทียบกับจำนวนนักเรียนทั้งหมด"),
        ("จำนวนนักเรียนระดับดีขึ้นไป (เกรด ≥ 3.0)", 
         "=COUNTIF('ตัดเกรด ห้อง 1'!M7:M24, \">=3.0\")", 
         "=COUNTIF('ตัดเกรด ห้อง 2'!M7:M22, \">=3.0\")", 
         "=COUNTIF('ตัดเกรด ห้อง 1'!M7:M24, \">=3.0\") + COUNTIF('ตัดเกรด ห้อง 2'!M7:M22, \">=3.0\")", 
         "เกรด 3.0, 3.5, 4.0"),
    ]
    
    for r_offset, r_data in enumerate(metrics, 2):
        cur_r = m_start + r_offset
        for c_offset, val in enumerate(r_data, 1):
            cell = ws.cell(row=cur_r, column=c_offset, value=val)
            cell.font = bold_data_font if c_offset == 1 else data_font
            cell.border = thin_border
            if c_offset == 1:
                cell.alignment = Alignment(horizontal='left', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='center', vertical='center')

    ov_widths = {1: 30, 2: 18, 3: 20, 4: 26, 5: 28, 6: 14, 7: 12}
    for c_idx, w in ov_widths.items():
        ws.column_dimensions[get_column_letter(c_idx)].width = w

create_overview_sheet()

wb.save(output_file)
print(f"Successfully generated complete grading simulation workbook: {output_file}")
