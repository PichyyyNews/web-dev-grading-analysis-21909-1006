import sys, io, os, shutil, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Roster definitions
ROOM1_ROSTER = [
    (1, '69219090001', '001', 'นายกฤศมน บุญเปี่ยม', 20, 3, 20, 4, 'เข้าสอบปกติ'),
    (2, '69219090002', '002', 'เด็กชายกษธณชล การรักเรียน', 9, 2, 14, 4, 'เข้าสอบปกติ'),
    (3, '69219090003', '003', 'นางสาวกัญญ์วรา ทองดี', 26, 30, 26, 18, 'เข้าสอบปกติ'),
    (4, '69219090004', '004', 'นายกันธรักษ์ พานทอง', 15, 5, 18, 5, 'เข้าสอบปกติ'),
    (5, '69219090005', '005', 'เด็กชายคุณานนต์ พิมพ์จันทร์', None, None, 27, 5, 'ขาดสอบกลางภาค'),
    (6, '69219090006', '006', 'เด็กชายเฉลิมวงศ์ มะลิใย', 15, 8, 22, 4, 'เข้าสอบปกติ'),
    (7, '69219090007', '007', 'นางสาวชลดา เรือนแก้ว', 15, 20, 15, 8.5, 'เข้าสอบปกติ'),
    (8, '69219090008', '008', 'นายณัฐกรณ์ กองทอง', None, None, None, None, 'ขาดสอบ'),
    (9, '69219090009', '009', 'นายณัฐพล ทวีทรัพย์', 13, 15, 17, 4, 'เข้าสอบปกติ'),
    (10, '69219090010', '010', 'นายต้นตระการ เพชรภู่', 11, 2, 6, 0, 'เข้าสอบปกติ'),
    (11, '69219090011', '011', 'นายตฤณ ตาดทอง', 21, 9, 25, 5, 'เข้าสอบปกติ'),
    (12, '69219090012', '012', 'นายธนพัฒน์ บุญเรือง', 20, 30, 22, 12, 'เข้าสอบปกติ'),
    (13, '69219090013', '013', 'นายธีรภัทร วาสิทธิ์เทพ', 16, 8, 17, 5, 'เข้าสอบปกติ'),
    (14, '69219090014', '014', 'นายภูเบต ควันทอง', None, None, None, None, 'ขาดสอบ'),
    (15, '69219090015', '015', 'นายภูวพิศิษฐ์ ใหลสกุล', 31, 40, 26, 19, 'เข้าสอบปกติ'),
    (16, '69219090016', '016', 'นางสาวมินตา บินอามัด', 16, 18, 10, 7, 'เข้าสอบปกติ'),
    (17, '69219090017', '017', 'นายอัรฮัม ยูนุ', 20, 15, 26, 5, 'เข้าสอบปกติ'),
    (18, '69219090018', '018', 'นายอานนท์ณัฏฐ์ ไวพจน์', 18, 15, 18, 8, 'เข้าสอบปกติ'),
]

ROOM2_ROSTER = [
    (1, '69219090020', '020', 'นายชุติวัต สุขมั่น', 8, 0, 15, 1, 'เข้าสอบปกติ'),
    (2, '69219090021', '021', 'เด็กชายณัพร รัตนสุวรรณ', 17, 5, 11, 4, 'เข้าสอบปกติ'),
    (3, '69219090022', '022', 'นายตุลยวัต ทาสีทอง', 19, 5, 19, 16, 'เข้าสอบปกติ'),
    (4, '69219090023', '023', 'เด็กชายธนกฤต แก้วตา', 23, 8, 22, 8, 'เข้าสอบปกติ'),
    (5, '69219090024', '024', 'นายธเนศพล กันเมือง', 12, 5, 16, 7, 'เข้าสอบปกติ'),
    (6, '69219090025', '025', 'นายนรากร จันทร์กล่ำ', None, None, None, None, 'ขาดสอบ'),
    (7, '69219090026', '026', 'นายนราวิชญ์ ทองระคนธ์', 19, 5, 19, 6, 'เข้าสอบปกติ'),
    (8, '69219090027', '027', 'นายนันทพงศ์ ช่ออุบล', 8, 0, 8, 0, 'เข้าสอบปกติ'),
    (9, '69219090028', '028', 'เด็กชายพงศ์พัทธ์ โสดา', 19, 5, 22, 15, 'เข้าสอบปกติ'),
    (10, '69219090029', '029', 'นายภูธเรศ อินทโชติ', 12, 5, 18, 13, 'เข้าสอบปกติ'),
    (11, '69219090030', '030', 'เด็กหญิงวราภรณ์ มั่นคง', 9, 0, 12, 0, 'เข้าสอบปกติ'),
    (12, '69219090031', '031', 'นายวีรวัฒน์ วงษ์ทิพย์', None, None, 13, 0, 'ขาดสอบกลางภาค'),
    (13, '69219090032', '032', 'นายวุฒิกร พาณิชย์กุลเกียรติ', 11, 7, 12, 2, 'เข้าสอบปกติ'),
    (14, '69219090033', '033', 'เด็กหญิงสโรชา มาเส็ง', 14, 0, 11, 0, 'เข้าสอบปกติ'),
    (15, '69219090034', '034', 'นายอธิป ทองดอนเถื่อน', 18, 7, 21, 8, 'เข้าสอบปกติ'),
    (16, '69219090035', '035', 'เด็กหญิงอารยา จันทน', 19, 5, 19, 8, 'เข้าสอบปกติ'),
]

# Extract assignment scores from Google Classroom files
def extract_assignments_r1():
    p = r'c:\Users\Newsk\Downloads\1ชทค\ห้อง1\คะแนน 21909-1006 พื้นฐานการสร้างเว็บไซต์  13-09-2569 (1).xlsx'
    wb = openpyxl.load_workbook(p)
    ws = wb.active
    data = {}
    for r in range(6, ws.max_row + 1):
        email = str(ws.cell(row=r, column=3).value or '').strip()
        m = re.search(r'69219090(\d{3})', email)
        if m:
            code = m.group(1)
            data[code] = [
                ws.cell(row=r, column=6).value,  # Ex1 เว็บไซต์แนะนำตัวเอง
                ws.cell(row=r, column=7).value,  # Ex2 เว็บไซต์ใบงาน
                ws.cell(row=r, column=14).value, # Ex3 CSS 3 รูปแบบ
                ws.cell(row=r, column=8).value,  # Ex4 colspan + rowspan
                ws.cell(row=r, column=9).value,  # Ex5 สร้าง form หน้าสมัครสมาชิก
                ws.cell(row=r, column=10).value, # Ex6 box model
                ws.cell(row=r, column=11).value, # Ex7 flex box layout
                ws.cell(row=r, column=12).value, # Ex8 Grid layout
                ws.cell(row=r, column=13).value, # โปรเจคเว็บจากดีไซน์ figma
                ws.cell(row=r, column=5).value,  # งานกลุ่ม
            ]
    return data

def extract_assignments_r2():
    p = r'c:\Users\Newsk\Downloads\1ชทค\ห้อง2\คะแนน 21909-1006 พื้นฐานการสร้างเว็บไซต์  13-09-2569.xlsx'
    wb = openpyxl.load_workbook(p)
    ws = wb.active
    data = {}
    for r in range(6, ws.max_row + 1):
        email = str(ws.cell(row=r, column=3).value or '').strip()
        m = re.search(r'69219090(\d{3})', email)
        if m:
            code = m.group(1)
            data[code] = [
                ws.cell(row=r, column=6).value,  # Ex1 เว็บไซต์แนะนำตัวเอง
                ws.cell(row=r, column=7).value,  # Ex2 เว็บไซต์ใบงาน
                ws.cell(row=r, column=12).value, # Ex3 CSS 3 รูปแบบ
                ws.cell(row=r, column=8).value,  # Ex4 colspan + rowspan
                ws.cell(row=r, column=9).value,  # Ex5 CSS boxmodel
                ws.cell(row=r, column=13).value, # การบ้านที่ 1 Form สมัครสมาชิก
                ws.cell(row=r, column=14).value, # Ex7 flex box layout
                ws.cell(row=r, column=10).value, # Ex8 Grid layout
                ws.cell(row=r, column=11).value, # โปรเจคเว็บจากดีไซน์ figma
                ws.cell(row=r, column=5).value,  # งานกลุ่ม
            ]
    return data

def build_room_workbook(room_num, roster, assignments_data, assignment_headers, target_paths):
    print(f"Building complete workbook for Room {room_num}...")
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)
    
    # Fonts & Styles
    font_main = "Leelawadee UI"
    title_font = Font(name=font_main, size=15, bold=True, color="0D47A1")
    subtitle_font = Font(name=font_main, size=11, color="546E7A")
    group_font = Font(name=font_main, size=11, bold=True, color="FFFFFF")
    col_header_font = Font(name=font_main, size=10, bold=True, color="1A202C")
    data_font = Font(name=font_main, size=10, color="2D3748")
    bold_data_font = Font(name=font_main, size=10, bold=True, color="1A202C")
    stat_label_font = Font(name=font_main, size=10, bold=True, color="1A365D")
    
    # Fills
    fill_info = PatternFill("solid", fgColor="1A365D")      # Dark Navy
    fill_assign = PatternFill("solid", fgColor="00695C")    # Deep Teal
    fill_mid = PatternFill("solid", fgColor="37474F")       # Slate Blue
    fill_fin = PatternFill("solid", fgColor="1565C0")       # Royal Blue
    fill_tot = PatternFill("solid", fgColor="4A148C")       # Deep Purple
    fill_col_head = PatternFill("solid", fgColor="F1F5F9")  # Very Light Gray
    fill_zebra = PatternFill("solid", fgColor="F8FAFC")     # Zebra subtle
    fill_written = PatternFill("solid", fgColor="E8F5E9")   # Soft Mint Green (written scores)
    fill_absent = PatternFill("solid", fgColor="FFEBEE")    # Soft Light Red
    fill_stat = PatternFill("solid", fgColor="E2E8F0")      # Stat header
    fill_highlight = PatternFill("solid", fgColor="FFF9C4") # Soft Yellow
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    double_bottom_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='double', color='1A365D')
    )

    # =========================================================================
    # SHEET 1: ตารางคะแนนรวมทั้งสิ้น (ครบทุกส่วน)
    # =========================================================================
    ws1 = wb.create_sheet(title='ตารางคะแนนรวม (เก็บ-กลาง-ปลาย)')
    ws1.views.sheetView[0].showGridLines = True
    
    # Banner
    ws1.merge_cells('A1:Y1')
    ws1['A1'] = f"📊 แบบบันทึกคะแนนผลการเรียนรู้ครบวงจร (คะแนนเก็บ + กลางภาค + ปลายภาค) - ห้อง {room_num}"
    ws1['A1'].font = title_font
    ws1['A1'].alignment = Alignment(horizontal='left', vertical='center')
    
    ws1.merge_cells('A2:Y2')
    ws1['A2'] = f"รหัสวิชา 21909-1006 พื้นฐานการสร้างเว็บไซต์  |  ระดับชั้น ปวช. 1 ชทค. ห้อง {room_num} (จำนวนนักเรียน {len(roster)} คน)  |  ภาคเรียนที่ 1 ปีการศึกษา 2569"
    ws1['A2'].font = subtitle_font
    ws1['A2'].alignment = Alignment(horizontal='left', vertical='center')
    
    # Row 4: Super Headers
    groups = [
        ('A4:D4', 'ข้อมูลนักศึกษา', fill_info),
        ('E4:O4', 'คะแนนเก็บแบบฝึกหัด (Coursework 100 คะแนน)', fill_assign),
        ('P4:R4', 'คะแนนสอบกลางภาค (Midterm)', fill_mid),
        ('S4:U4', 'คะแนนสอบปลายภาค (Final Examination)', fill_fin),
        ('V4:Y4', 'สรุปผลคะแนนรวมทั้งภาคเรียน', fill_tot)
    ]
    for cell_rng, g_name, g_fill in groups:
        ws1.merge_cells(cell_rng)
        top_left = cell_rng.split(':')[0]
        ws1[top_left] = g_name
        ws1[top_left].font = group_font
        ws1[top_left].alignment = Alignment(horizontal='center', vertical='center')
        
        # Apply fill to all cells in merge range
        start_col, start_row = openpyxl.utils.coordinate_to_tuple(cell_rng.split(':')[0])
        end_col, end_row = openpyxl.utils.coordinate_to_tuple(cell_rng.split(':')[1])
        for r_idx in range(start_row, end_row + 1):
            for c_idx in range(start_col, end_col + 1):
                c = ws1.cell(row=r_idx, column=c_idx)
                c.fill = g_fill
                c.border = thin_border
                
    # Row 5: Column Subheaders
    sub_headers = [
        "เลขที่", "รหัสนักศึกษา", "รหัส\n(3 ตัว)", "ชื่อ - นามสกุล",
        # E - N: 10 assignments
        assignment_headers[0], assignment_headers[1], assignment_headers[2], assignment_headers[3],
        assignment_headers[4], assignment_headers[5], assignment_headers[6], assignment_headers[7],
        assignment_headers[8], assignment_headers[9],
        "รวมคะแนนเก็บ\n(เต็ม 100)",
        # P - R: Midterm
        "ข้อกา\n(ปรนัย)", "ข้อเขียน\n(อัตนัย)", "รวม\nกลางภาค",
        # S - U: Final
        "ข้อกา (30)\n(ปรนัย)", "ข้อเขียน (20)\n(อัตนัย)", "รวมปลายภาค\n(เต็ม 50)",
        # V - Y: Grand Totals
        "รวมคะแนนสอบ\n(กลาง+ปลาย)", "รวมคะแนนสุทธิ\n(เก็บ+สอบ)", "ลำดับที่\n(Rank)", "สถานะ / หมายเหตุ"
    ]
    
    ws1.row_dimensions[4].height = 28
    ws1.row_dimensions[5].height = 36
    ws1.row_dimensions[6].height = 24
    
    for c_idx, h_text in enumerate(sub_headers, 1):
        c = ws1.cell(row=5, column=c_idx, value=h_text)
        c.font = col_header_font
        c.fill = fill_col_head
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = thin_border
        
    # Row 6: Max Points
    max_row_vals = [
        None, None, None, "คะแนนเต็ม",
        10, 10, 10, 10, 10, 10, 10, 10, 10, 10, 100,
        "-", "-", "-",
        30, 20, 50,
        "-", "-", "-", "-"
    ]
    for c_idx, val in enumerate(max_row_vals, 1):
        c = ws1.cell(row=6, column=c_idx, value=val)
        c.font = bold_data_font
        c.fill = PatternFill("solid", fgColor="E0F2FE") # Soft ice blue
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = thin_border
        if c_idx == 4:
            c.alignment = Alignment(horizontal='right', vertical='center')
            
    # Student Data Rows (7 to 6 + len(roster))
    start_r = 7
    for idx, (no, full_id, code, name, mid_obj, mid_wri, fin_obj, fin_wri, status) in enumerate(roster):
        cur_r = start_r + idx
        ws1.row_dimensions[cur_r].height = 22
        is_even = (idx % 2 == 1)
        row_fill = fill_zebra if is_even else PatternFill("solid", fgColor="FFFFFF")
        
        # Student info
        ws1.cell(row=cur_r, column=1, value=no).alignment = Alignment(horizontal='center', vertical='center')
        ws1.cell(row=cur_r, column=2, value=str(full_id)).alignment = Alignment(horizontal='center', vertical='center')
        ws1.cell(row=cur_r, column=3, value=str(code)).alignment = Alignment(horizontal='center', vertical='center')
        ws1.cell(row=cur_r, column=4, value=name).alignment = Alignment(horizontal='left', vertical='center')
        
        for c in range(1, 5):
            ws1.cell(row=cur_r, column=c).font = data_font
            ws1.cell(row=cur_r, column=c).fill = row_fill
            ws1.cell(row=cur_r, column=c).border = thin_border
            
        # Assignments E - N
        as_scores = assignments_data.get(code, [None]*10)
        for a_idx, sc_val in enumerate(as_scores):
            col_idx = 5 + a_idx
            c = ws1.cell(row=cur_r, column=col_idx, value=sc_val)
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.font = data_font
            c.fill = row_fill
            c.border = thin_border
            if sc_val is not None:
                c.number_format = '0' if isinstance(sc_val, int) or (isinstance(sc_val, float) and sc_val.is_integer()) else '0.0'
                
        # Assignment Total (Col 15 / O)
        c_tot_as = ws1.cell(row=cur_r, column=15)
        c_tot_as.value = f'=IF(COUNT(E{cur_r}:N{cur_r})>0, SUM(E{cur_r}:N{cur_r}), "-")'
        c_tot_as.alignment = Alignment(horizontal='center', vertical='center')
        c_tot_as.font = bold_data_font
        c_tot_as.fill = PatternFill("solid", fgColor="E0F2F1") # Soft teal tint
        c_tot_as.border = thin_border
        
        # Midterm P - R
        ws1.cell(row=cur_r, column=16, value=mid_obj).alignment = Alignment(horizontal='center', vertical='center')
        ws1.cell(row=cur_r, column=17, value=mid_wri).alignment = Alignment(horizontal='center', vertical='center')
        c_mid_tot = ws1.cell(row=cur_r, column=18)
        c_mid_tot.value = f'=IF(COUNT(P{cur_r}:Q{cur_r})>0, SUM(P{cur_r}:Q{cur_r}), "-")'
        c_mid_tot.alignment = Alignment(horizontal='center', vertical='center')
        for col_idx in [16, 17, 18]:
            c = ws1.cell(row=cur_r, column=col_idx)
            c.font = bold_data_font if col_idx == 18 else data_font
            c.fill = row_fill
            c.border = thin_border
            
        # Final S - U
        c_fin_obj = ws1.cell(row=cur_r, column=19, value=fin_obj)
        c_fin_obj.alignment = Alignment(horizontal='center', vertical='center')
        c_fin_obj.font = data_font
        c_fin_obj.fill = row_fill
        c_fin_obj.border = thin_border
        
        c_fin_wri = ws1.cell(row=cur_r, column=20, value=fin_wri)
        c_fin_wri.alignment = Alignment(horizontal='center', vertical='center')
        c_fin_wri.border = thin_border
        if fin_wri is not None:
            c_fin_wri.font = Font(name=font_main, size=10, bold=True, color="1B5E20")
            c_fin_wri.fill = fill_written
            c_fin_wri.number_format = '0.0' if isinstance(fin_wri, float) and not fin_wri.is_integer() else '0'
        else:
            c_fin_wri.font = Font(name=font_main, size=10, color="C62828")
            c_fin_wri.fill = fill_absent
            
        c_fin_tot = ws1.cell(row=cur_r, column=21)
        c_fin_tot.value = f'=IF(COUNT(S{cur_r}:T{cur_r})>0, SUM(S{cur_r}:T{cur_r}), "-")'
        c_fin_tot.alignment = Alignment(horizontal='center', vertical='center')
        c_fin_tot.font = bold_data_font
        c_fin_tot.fill = PatternFill("solid", fgColor="E3F2FD") # Soft blue tint
        c_fin_tot.border = thin_border
        
        # Grand Totals V - Y
        # V: Total Exam (Mid + Fin)
        c_exam_tot = ws1.cell(row=cur_r, column=22)
        c_exam_tot.value = f'=IF(OR(ISNUMBER(R{cur_r}), ISNUMBER(U{cur_r})), IF(ISNUMBER(R{cur_r}), R{cur_r}, 0) + IF(ISNUMBER(U{cur_r}), U{cur_r}, 0), "-")'
        c_exam_tot.alignment = Alignment(horizontal='center', vertical='center')
        c_exam_tot.font = bold_data_font
        c_exam_tot.fill = row_fill
        c_exam_tot.border = thin_border
        
        # W: Grand Total (Coursework + Mid + Fin)
        c_grand = ws1.cell(row=cur_r, column=23)
        c_grand.value = f'=IF(OR(ISNUMBER(O{cur_r}), ISNUMBER(V{cur_r})), IF(ISNUMBER(O{cur_r}), O{cur_r}, 0) + IF(ISNUMBER(V{cur_r}), V{cur_r}, 0), "-")'
        c_grand.alignment = Alignment(horizontal='center', vertical='center')
        c_grand.font = Font(name=font_main, size=10, bold=True, color="4A148C")
        c_grand.fill = PatternFill("solid", fgColor="F3E5F5") # Soft purple tint
        c_grand.border = thin_border
        
        # X: Rank (based on Grand Total W)
        last_stu_r = start_r + len(roster) - 1
        c_rank = ws1.cell(row=cur_r, column=24)
        c_rank.value = f'=IF(ISNUMBER(W{cur_r}), RANK(W{cur_r}, $W$7:$W${last_stu_r}), "-")'
        c_rank.alignment = Alignment(horizontal='center', vertical='center')
        c_rank.font = bold_data_font
        c_rank.fill = row_fill
        c_rank.border = thin_border
        
        # Y: Status
        c_stat = ws1.cell(row=cur_r, column=25, value=status)
        c_stat.alignment = Alignment(horizontal='center', vertical='center')
        c_stat.fill = row_fill
        c_stat.border = thin_border
        if "ขาด" in status:
            c_stat.font = Font(name=font_main, size=10, bold=True, color="C62828")
        else:
            c_stat.font = Font(name=font_main, size=10, color="2E7D32")
            
    # Summary Rows (Average, Max, Min, S.D.)
    stat_rows_def = [
        ("คะแนนเฉลี่ย (Average)", "AVERAGE", True),
        ("คะแนนสูงสุด (Max)", "MAX", False),
        ("คะแนนต่ำสุด (Min)", "MIN", False),
        ("ส่วนเบี่ยงเบนมาตรฐาน (S.D.)", "STDEV.P", True),
    ]
    
    last_r = start_r + len(roster) - 1
    for s_idx, (s_label, s_func, is_round) in enumerate(stat_rows_def):
        r_num = last_r + 1 + s_idx
        ws1.row_dimensions[r_num].height = 22
        
        # Merge A:D for label
        ws1.merge_cells(start_row=r_num, start_column=1, end_row=r_num, end_column=4)
        lbl_cell = ws1.cell(row=r_num, column=1, value=s_label)
        lbl_cell.font = stat_label_font
        lbl_cell.alignment = Alignment(horizontal='right', vertical='center')
        
        for c in range(1, 5):
            ws1.cell(row=r_num, column=c).fill = fill_stat
            ws1.cell(row=r_num, column=c).border = thin_border if s_idx < 3 else double_bottom_border
            
        # For columns E to W (5 to 23)
        for col_idx in range(5, 24):
            c_letter = get_column_letter(col_idx)
            c = ws1.cell(row=r_num, column=col_idx)
            if is_round:
                c.value = f'=IFERROR(ROUND({s_func}({c_letter}7:{c_letter}{last_r}), 2), "-")'
            else:
                c.value = f'=IFERROR({s_func}({c_letter}7:{c_letter}{last_r}), "-")'
            c.font = bold_data_font
            c.fill = fill_stat
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.border = thin_border if s_idx < 3 else double_bottom_border
            
        # X & Y blank border
        for c_idx in [24, 25]:
            c = ws1.cell(row=r_num, column=c_idx)
            c.fill = fill_stat
            c.border = thin_border if s_idx < 3 else double_bottom_border

    # Set Column widths
    col_widths = {
        1: 7,   # No.
        2: 15,  # Student ID
        3: 9,   # 3-digit Code
        4: 26,  # Name
        # E - N: Assignments
        5: 12, 6: 12, 7: 12, 8: 12, 9: 13, 10: 12, 11: 12, 12: 12, 13: 13, 14: 11,
        15: 14, # Total Assignment
        # P - R: Midterm
        16: 11, 17: 11, 18: 12,
        # S - U: Final
        19: 12, 20: 13, 21: 14,
        # V - Y: Totals
        22: 15, 23: 16, 24: 11, 25: 16
    }
    for c_idx, w in col_widths.items():
        ws1.column_dimensions[get_column_letter(c_idx)].width = w

    # =========================================================================
    # SHEET 2: คะแนนเก็บแบบฝึกหัด (แจกแจงละเอียด 10 ชิ้นงาน)
    # =========================================================================
    ws2 = wb.create_sheet(title='คะแนนเก็บแบบฝึกหัด (แจกแจง)')
    ws2.views.sheetView[0].showGridLines = True
    
    ws2.merge_cells('A1:Q1')
    ws2['A1'] = f"📝 แบบบันทึกคะแนนเก็บแบบฝึกหัดและชิ้นงาน (เรียงลำดับตามหลักสูตร) - ห้อง {room_num}"
    ws2['A1'].font = title_font
    ws2['A1'].alignment = Alignment(horizontal='left', vertical='center')
    
    ws2.merge_cells('A2:Q2')
    ws2['A2'] = f"วิชา 21909-1006 พื้นฐานการสร้างเว็บไซต์  |  แบบฝึกหัดที่ 1 - 8 + โปรเจค Figma + งานกลุ่ม  |  เต็ม 100 คะแนน"
    ws2['A2'].font = subtitle_font
    ws2['A2'].alignment = Alignment(horizontal='left', vertical='center')
    
    # Headers
    as_full_headers = [
        "เลขที่", "รหัสนักศึกษา", "รหัส", "ชื่อ - นามสกุล",
        assignment_headers[0], assignment_headers[1], assignment_headers[2], assignment_headers[3],
        assignment_headers[4], assignment_headers[5], assignment_headers[6], assignment_headers[7],
        assignment_headers[8], assignment_headers[9],
        "รวมคะแนนเก็บ\n(เต็ม 100)", "ร้อยละ\n(%)", "ลำดับที่\n(Rank)"
    ]
    ws2.row_dimensions[4].height = 36
    ws2.row_dimensions[5].height = 24
    
    for c_idx, h_text in enumerate(as_full_headers, 1):
        c = ws2.cell(row=4, column=c_idx, value=h_text)
        c.font = Font(name=font_main, size=10, bold=True, color="FFFFFF")
        c.fill = fill_assign
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = thin_border
        
    # Row 5: Max Points
    as_max_vals = [None, None, None, "คะแนนเต็ม", 10, 10, 10, 10, 10, 10, 10, 10, 10, 10, 100, "100.0%", "-"]
    for c_idx, val in enumerate(as_max_vals, 1):
        c = ws2.cell(row=5, column=c_idx, value=val)
        c.font = bold_data_font
        c.fill = PatternFill("solid", fgColor="E0F2F1")
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = thin_border
        if c_idx == 4:
            c.alignment = Alignment(horizontal='right', vertical='center')
            
    # Student rows
    for idx, (no, full_id, code, name, *_) in enumerate(roster):
        cur_r = 6 + idx
        ws2.row_dimensions[cur_r].height = 22
        is_even = (idx % 2 == 1)
        row_fill = fill_zebra if is_even else PatternFill("solid", fgColor="FFFFFF")
        
        ws2.cell(row=cur_r, column=1, value=no).alignment = Alignment(horizontal='center', vertical='center')
        ws2.cell(row=cur_r, column=2, value=str(full_id)).alignment = Alignment(horizontal='center', vertical='center')
        ws2.cell(row=cur_r, column=3, value=str(code)).alignment = Alignment(horizontal='center', vertical='center')
        ws2.cell(row=cur_r, column=4, value=name).alignment = Alignment(horizontal='left', vertical='center')
        
        for c in range(1, 5):
            ws2.cell(row=cur_r, column=c).font = data_font
            ws2.cell(row=cur_r, column=c).fill = row_fill
            ws2.cell(row=cur_r, column=c).border = thin_border
            
        as_scores = assignments_data.get(code, [None]*10)
        for a_idx, sc_val in enumerate(as_scores):
            col_idx = 5 + a_idx
            c = ws2.cell(row=cur_r, column=col_idx, value=sc_val)
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.font = data_font
            c.fill = row_fill
            c.border = thin_border
            if sc_val is not None:
                c.number_format = '0' if isinstance(sc_val, int) or (isinstance(sc_val, float) and sc_val.is_integer()) else '0.0'
                
        # Total
        c_tot = ws2.cell(row=cur_r, column=15)
        c_tot.value = f'=IF(COUNT(E{cur_r}:N{cur_r})>0, SUM(E{cur_r}:N{cur_r}), "-")'
        c_tot.alignment = Alignment(horizontal='center', vertical='center')
        c_tot.font = bold_data_font
        c_tot.fill = PatternFill("solid", fgColor="E0F2F1")
        c_tot.border = thin_border
        
        # Percentage
        c_pct = ws2.cell(row=cur_r, column=16)
        c_pct.value = f'=IF(ISNUMBER(O{cur_r}), ROUND(O{cur_r}/$O$5*100, 1) & "%", "-")'
        c_pct.alignment = Alignment(horizontal='center', vertical='center')
        c_pct.font = data_font
        c_pct.fill = row_fill
        c_pct.border = thin_border
        
        # Rank
        last_as_r = 6 + len(roster) - 1
        c_rk = ws2.cell(row=cur_r, column=17)
        c_rk.value = f'=IF(ISNUMBER(O{cur_r}), RANK(O{cur_r}, $O$6:$O${last_as_r}), "-")'
        c_rk.alignment = Alignment(horizontal='center', vertical='center')
        c_rk.font = bold_data_font
        c_rk.fill = row_fill
        c_rk.border = thin_border
        
    # Summary stats for Sheet 2
    last_as_r = 6 + len(roster) - 1
    for s_idx, (s_label, s_func, is_round) in enumerate(stat_rows_def):
        r_num = last_as_r + 1 + s_idx
        ws2.row_dimensions[r_num].height = 22
        ws2.merge_cells(start_row=r_num, start_column=1, end_row=r_num, end_column=4)
        lbl_cell = ws2.cell(row=r_num, column=1, value=s_label)
        lbl_cell.font = stat_label_font
        lbl_cell.alignment = Alignment(horizontal='right', vertical='center')
        
        for c in range(1, 5):
            ws2.cell(row=r_num, column=c).fill = fill_stat
            ws2.cell(row=r_num, column=c).border = thin_border if s_idx < 3 else double_bottom_border
            
        for col_idx in range(5, 16):
            c_letter = get_column_letter(col_idx)
            c = ws2.cell(row=r_num, column=col_idx)
            if is_round:
                c.value = f'=IFERROR(ROUND({s_func}({c_letter}6:{c_letter}{last_as_r}), 2), "-")'
            else:
                c.value = f'=IFERROR({s_func}({c_letter}6:{c_letter}{last_as_r}), "-")'
            c.font = bold_data_font
            c.fill = fill_stat
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.border = thin_border if s_idx < 3 else double_bottom_border
            
        for c_idx in [16, 17]:
            c = ws2.cell(row=r_num, column=c_idx)
            c.fill = fill_stat
            c.border = thin_border if s_idx < 3 else double_bottom_border
            
    as_widths = {1: 7, 2: 15, 3: 8, 4: 26, 5: 13, 6: 13, 7: 13, 8: 13, 9: 13, 10: 13, 11: 13, 12: 13, 13: 13, 14: 11, 15: 15, 16: 11, 17: 11}
    for c_idx, w in as_widths.items():
        ws2.column_dimensions[get_column_letter(c_idx)].width = w

    # Copy Sheet 3 (ตรวจปลายภาคละเอียด) & Sheet 4 (เฉลย) from existing files if available
    # We can preserve or re-create them cleanly
    template_path = r'c:\Users\Newsk\Downloads\1ชทค\คะแนนสอบปลายภาค_21909-1006_พื้นฐานการสร้างเว็บไซต์.xlsx'
    if os.path.exists(template_path):
        wb_master = openpyxl.load_workbook(template_path)
        sheet_mc_name = f'ห้อง {room_num} - ตรวจข้อสอบละเอียด'
        if sheet_mc_name in wb_master.sheetnames:
            ws_src = wb_master[sheet_mc_name]
            ws_target = wb.create_sheet(title='ตรวจปลายภาคละเอียด (30 ข้อ)')
            ws_target.views.sheetView[0].showGridLines = True
            for r in range(1, ws_src.max_row + 1):
                ws_target.row_dimensions[r].height = ws_src.row_dimensions[r].height
                for c in range(1, ws_src.max_column + 1):
                    src_cell = ws_src.cell(row=r, column=c)
                    dst_cell = ws_target.cell(row=r, column=c, value=src_cell.value)
                    if src_cell.has_style:
                        dst_cell.font = Font(name=src_cell.font.name, size=src_cell.font.size, bold=src_cell.font.bold, color=src_cell.font.color)
                        dst_cell.fill = PatternFill(src_cell.fill.fill_type, fgColor=src_cell.fill.fgColor)
                        dst_cell.alignment = Alignment(horizontal=src_cell.alignment.horizontal, vertical=src_cell.alignment.vertical, wrap_text=src_cell.alignment.wrap_text)
                        dst_cell.border = thin_border
            for col_l in ws_src.column_dimensions:
                ws_target.column_dimensions[col_l].width = ws_src.column_dimensions[col_l].width

        if 'เฉลยและเกณฑ์การให้คะแนน' in wb_master.sheetnames:
            ws_key = wb_master['เฉลยและเกณฑ์การให้คะแนน']
            ws_key_target = wb.create_sheet(title='เฉลยและเกณฑ์การให้คะแนน')
            ws_key_target.views.sheetView[0].showGridLines = True
            for r in range(1, ws_key.max_row + 1):
                ws_key_target.row_dimensions[r].height = ws_key.row_dimensions[r].height
                for c in range(1, ws_key.max_column + 1):
                    src_cell = ws_key.cell(row=r, column=c)
                    dst_cell = ws_key_target.cell(row=r, column=c, value=src_cell.value)
                    if src_cell.has_style:
                        dst_cell.font = Font(name=src_cell.font.name, size=src_cell.font.size, bold=src_cell.font.bold, color=src_cell.font.color)
                        dst_cell.fill = PatternFill(src_cell.fill.fill_type, fgColor=src_cell.fill.fgColor)
                        dst_cell.alignment = Alignment(horizontal=src_cell.alignment.horizontal, vertical=src_cell.alignment.vertical, wrap_text=src_cell.alignment.wrap_text)
                        dst_cell.border = thin_border
            for col_l in ws_key.column_dimensions:
                ws_key_target.column_dimensions[col_l].width = ws_key.column_dimensions[col_l].width

    # Save to all target paths
    for tp in target_paths:
        wb.save(tp)
        print(f"Saved: {tp}")

# Prepare Headers for each room
r1_as_headers = [
    "แบบฝึกหัดที่ 1\nแนะนำตัวเอง (10)",
    "แบบฝึกหัดที่ 2\nใบงาน (10)",
    "แบบฝึกหัดที่ 3\nCSS 3 รูปแบบ (10)",
    "แบบฝึกหัดที่ 4\ncolspan+rowspan (10)",
    "แบบฝึกหัดที่ 5\nForm สมัครสมาชิก (10)",
    "แบบฝึกหัดที่ 6\nBox Model (10)",
    "แบบฝึกหัดที่ 7\nFlexbox Layout (10)",
    "แบบฝึกหัดที่ 8\nGrid Layout (10)",
    "โปรเจคเว็บ\nFigma Design (10)",
    "งานกลุ่ม\n(10)"
]

r2_as_headers = [
    "แบบฝึกหัดที่ 1\nแนะนำตัวเอง (10)",
    "แบบฝึกหัดที่ 2\nใบงาน (10)",
    "แบบฝึกหัดที่ 3\nCSS 3 รูปแบบ (10)",
    "แบบฝึกหัดที่ 4\ncolspan+rowspan (10)",
    "แบบฝึกหัดที่ 5\nCSS Boxmodel (10)",
    "การบ้านที่ 1\nForm สมัครสมาชิก (10)",
    "แบบฝึกหัดที่ 7\nFlexbox Layout (10)",
    "แบบฝึกหัดที่ 8\nGrid Layout (10)",
    "โปรเจคเว็บ\nFigma Design (10)",
    "งานกลุ่ม\n(10)"
]

if __name__ == '__main__':
    as_r1_data = extract_assignments_r1()
    as_r2_data = extract_assignments_r2()

    # Target paths
    r1_targets = [
        r'c:\Users\Newsk\Downloads\1ชทค\คะแนนรวม_21909-1006_พื้นฐานการสร้างเว็บไซต์_ห้อง1.xlsx',
        r'c:\Users\Newsk\Downloads\1ชทค\ห้อง1\คะแนนรวม_21909-1006_พื้นฐานการสร้างเว็บไซต์_ห้อง1.xlsx'
    ]
    r2_targets = [
        r'c:\Users\Newsk\Downloads\1ชทค\คะแนนรวม_21909-1006_พื้นฐานการสร้างเว็บไซต์_ห้อง2.xlsx',
        r'c:\Users\Newsk\Downloads\1ชทค\ห้อง2\คะแนนรวม_21909-1006_พื้นฐานการสร้างเว็บไซต์_ห้อง2.xlsx'
    ]

    build_room_workbook(1, ROOM1_ROSTER, as_r1_data, r1_as_headers, r1_targets)
    build_room_workbook(2, ROOM2_ROSTER, as_r2_data, r2_as_headers, r2_targets)

    print("Both classroom workbooks generated successfully!")

